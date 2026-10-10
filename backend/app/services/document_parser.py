import os
import io
import logging
from typing import List, Dict, Any
import pypdf
import docx

logger = logging.getLogger("backend.parser")


class DocumentParser:
    @staticmethod
    def parse_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
        pages = []
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                text = text.strip()
                if text:
                    pages.append({
                        "page_num": i + 1,
                        "text": text
                    })
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

        if not pages:
            # Fallback if no text extracted (e.g. empty PDF or image-only)
            pages.append({"page_num": 1, "text": "[Empty or unscannable PDF content]"})
        return pages

    @staticmethod
    def parse_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
        pages = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n\n".join(paragraphs)
            
            # Approximate pages based on character length (~2000 chars per page)
            approx_page_size = 2000
            if len(full_text) <= approx_page_size:
                pages.append({"page_num": 1, "text": full_text})
            else:
                chunks = [full_text[i:i + approx_page_size] for i in range(0, len(full_text), approx_page_size)]
                for i, chunk_text in enumerate(chunks):
                    pages.append({"page_num": i + 1, "text": chunk_text.strip()})
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

        if not pages:
            pages.append({"page_num": 1, "text": "[Empty DOCX content]"})
        return pages

    @staticmethod
    def parse_txt(file_bytes: bytes) -> List[Dict[str, Any]]:
        encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]
        text = ""
        for enc in encodings:
            try:
                text = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        text = text.strip()
        pages = []
        approx_page_size = 2000
        if len(text) <= approx_page_size:
            pages.append({"page_num": 1, "text": text or "[Empty text content]"})
        else:
            chunks = [text[i:i + approx_page_size] for i in range(0, len(text), approx_page_size)]
            for i, chunk_text in enumerate(chunks):
                pages.append({"page_num": i + 1, "text": chunk_text.strip()})
        return pages

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> List[Dict[str, Any]]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_bytes)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_bytes)
        elif ext in [".txt", ".md", ".csv", ".json", ".log"]:
            return cls.parse_txt(file_bytes)
        else:
            # Try parsing as plain text by default
            return cls.parse_txt(file_bytes)
