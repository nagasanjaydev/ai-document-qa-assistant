import asyncio
import io
import os
import sys

# Ensure root in sys.path
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from backend.app.database import db_manager, get_users_collection, get_documents_collection, get_chunks_collection
from backend.app.services.document_parser import DocumentParser
from backend.app.services.chunker import TextChunker
from backend.app.services.embedding_service import EmbeddingService
from backend.app.services.rag_service import RAGService
from backend.app.services.chat_service import ChatService
from backend.app.services.auth_service import hash_password, verify_password, create_access_token, decode_token
import docx
import pypdf


async def run_tests():
    print("=== Testing Database Connection ===")
    await db_manager.connect()
    print("Database status:", db_manager.status_message)
    print("Is real MongoDB connected:", db_manager.is_mongodb_connected)

    print("\n=== Testing Auth & Password Hashing ===")
    pw = "SecretPassword123!"
    hashed = hash_password(pw)
    assert verify_password(pw, hashed), "Password verification failed"
    token = create_access_token({"sub": "test-user-id", "username": "alice"})
    payload = decode_token(token)
    assert payload["sub"] == "test-user-id", "JWT decode failed"
    print("Auth & JWT verified successfully!")

    print("\n=== Testing Document Parsing ===")
    # 1. Text Parsing
    sample_text = (
        "Artificial Intelligence in Healthcare.\n\n"
        "Machine learning models are rapidly transforming clinical workflows, "
        "allowing medical practitioners to diagnose complex diseases earlier. "
        "Natural language processing helps extract insights from unstructured electronic health records.\n\n"
        "Key Challenges:\n"
        "Patient data privacy, algorithmic bias, and regulatory compliance under HIPAA are critical concerns."
    )
    txt_pages = DocumentParser.parse_file("healthcare_ai.txt", sample_text.encode("utf-8"))
    assert len(txt_pages) > 0, "TXT parsing failed"
    print(f"TXT parsed: {len(txt_pages)} page(s)")

    # 2. DOCX Generation & Parsing
    doc = docx.Document()
    doc.add_heading("Quantum Computing Advances", 0)
    doc.add_paragraph("Superconducting qubits have demonstrated high coherence times.")
    doc.add_paragraph("Quantum error correction is essential for fault-tolerant universal quantum computing.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()
    docx_pages = DocumentParser.parse_file("quantum.docx", docx_bytes)
    assert len(docx_pages) > 0, "DOCX parsing failed"
    print(f"DOCX parsed: {len(docx_pages)} page(s), sample: {docx_pages[0]['text'][:40]}...")

    print("\n=== Testing Chunking ===")
    chunker = TextChunker(chunk_size=150, chunk_overlap=30)
    chunks = chunker.chunk_document_pages(txt_pages)
    assert len(chunks) > 0, "Chunking produced no chunks"
    print(f"Generated {len(chunks)} chunks with sliding window:")
    for c in chunks:
        print(f"  Chunk #{c['chunk_index']} (len {c['char_count']}): {c['content'][:50]}...")

    print("\n=== Testing Embeddings & Cosine Similarity ===")
    emb1 = await EmbeddingService.get_embedding("Machine learning for clinical diagnosis")
    emb2 = await EmbeddingService.get_embedding("Healthcare artificial intelligence and medicine")
    emb3 = await EmbeddingService.get_embedding("Planetary exploration of Jupiter's icy moons")

    sim_related = EmbeddingService.cosine_similarity(emb1, emb2)
    sim_unrelated = EmbeddingService.cosine_similarity(emb1, emb3)
    print(f"Similarity (ML Diagnosis vs Healthcare AI): {sim_related:.4f}")
    print(f"Similarity (ML Diagnosis vs Jupiter Moons): {sim_unrelated:.4f}")
    assert sim_related > sim_unrelated, "Semantic similarity relationship failed"

    print("\n=== Testing Mock RAG Indexing and Query ===")
    test_user_id = "test-user-id"
    test_doc_id = "test-doc-id"
    chunks_col = get_chunks_collection()
    docs_col = get_documents_collection()

    await docs_col.insert_one({
        "id": test_doc_id,
        "_id": test_doc_id,
        "user_id": test_user_id,
        "filename": "healthcare_ai.txt",
        "file_type": "txt",
        "file_size": len(sample_text),
        "chunk_count": len(chunks),
        "summary": chunks[0]["content"][:100],
        "created_at": "2026-10-10T00:00:00Z"
    })

    chunk_records = []
    for c in chunks:
        chunk_emb = await EmbeddingService.get_embedding(c["content"])
        chunk_records.append({
            "id": f"chunk-{c['chunk_index']}",
            "_id": f"chunk-{c['chunk_index']}",
            "doc_id": test_doc_id,
            "user_id": test_user_id,
            "chunk_index": c["chunk_index"],
            "content": c["content"],
            "page_num": c["page_num"],
            "embedding": chunk_emb
        })
    await chunks_col.insert_many(chunk_records)

    # Perform retrieval
    query = "What are the main privacy challenges in clinical AI?"
    retrieved = await RAGService.retrieve_chunks(query, test_user_id, doc_id=test_doc_id, top_k=2)
    assert len(retrieved) > 0, "RAG retrieval returned no results"
    print(f"RAG retrieved top chunk score: {retrieved[0]['score']}, content: {retrieved[0]['content'][:60]}...")

    # Generate answer
    rag_resp = await RAGService.generate_grounded_response(query, retrieved)
    print(f"RAG Model used: {rag_resp['model_used']}")
    print(f"Citations count: {len(rag_resp['citations'])}")
    print(f"Answer preview: {rag_resp['answer'][:150]}...")

    # Chat session persistence
    session = await ChatService.get_or_create_session(test_user_id, doc_id=test_doc_id, title="Healthcare AI Chat")
    session_id = session["id"]
    await ChatService.add_message(session_id, test_user_id, "user", query)
    await ChatService.add_message(session_id, test_user_id, "assistant", rag_resp["answer"], [c.model_dump() for c in rag_resp["citations"]])
    msgs = await ChatService.get_session_messages(session_id, test_user_id)
    assert len(msgs) == 2, f"Expected 2 messages, got {len(msgs)}"
    print(f"Chat persistence verified: {len(msgs)} messages stored.")

    await db_manager.close()
    print("\nALL BACKEND TESTS PASSED SUCCESSFULLY! ✅")


if __name__ == "__main__":
    asyncio.run(run_tests())
